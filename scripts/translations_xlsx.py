"""Builds the Excel translation workbook, and reads the finished one back.

    python3 scripts/translations_xlsx.py            # translation/*.json -> .xlsx
    python3 scripts/translations_xlsx.py --merge    # .xlsx -> translation/*.json

Needs openpyxl (`pip3 install openpyxl`). It is not a project dependency —
nothing in the app imports it, and the JSON and CSV paths in
`scripts/translations-to-sheets.mjs` cover the same ground without it. This
exists because one workbook with a tab per area is the nicest thing to hand a
translator, and because "save all 15 sheets as CSV" is a chore nobody should do
by hand.

The JSON under `translation/` stays the source of truth throughout. This is a
view onto it, the same way the CSVs are.
"""

import json
import pathlib
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
JSON_DIR = ROOT / "translation"
WORKBOOK = JSON_DIR / "georgian-translation.xlsx"

COLUMNS = ["key", "english", "georgian", "source"]
GEORGIAN_COLUMN = 3

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
DONE_FILL = PatternFill("solid", fgColor="E2EFDA")


def namespaces():
    return sorted(p.stem for p in JSON_DIR.glob("*.json"))


def load(namespace):
    return json.loads((JSON_DIR / f"{namespace}.json").read_text(encoding="utf-8"))


def write_cell(sheet, row, column, value):
    """Writes text, never a formula.

    Excel treats a string beginning with `=` as a formula, and a handful of
    source strings legitimately start with one. Forcing the cell's type to
    string keeps those intact instead of turning them into `#NAME?`.
    """
    cell = sheet.cell(row=row, column=column, value=value)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        cell.data_type = "s"
    return cell


def build():
    book = Workbook()
    book.remove(book.active)

    overview = book.create_sheet("Overview")
    counts = {}

    for namespace in namespaces():
        entries = load(namespace)
        sheet = book.create_sheet(namespace)

        for index, name in enumerate(COLUMNS, start=1):
            cell = write_cell(sheet, 1, index, name)
            cell.font = Font(name=FONT, bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(vertical="center")

        # Untranslated first, so the work sits at the top of every tab and the
        # remaining block visibly shrinks as it is done.
        rows = sorted(entries.items(), key=lambda kv: (bool(kv[1].get("ka")), kv[0]))

        for offset, (key, entry) in enumerate(rows, start=2):
            georgian = entry.get("ka") or ""
            values = [key, entry.get("en", ""), georgian, entry.get("_source", "")]

            for index, value in enumerate(values, start=1):
                cell = write_cell(sheet, offset, index, value)
                cell.font = Font(name=FONT)
                cell.alignment = Alignment(
                    vertical="top", wrap_text=index in (2, 3)
                )
                if index == GEORGIAN_COLUMN:
                    cell.fill = DONE_FILL if georgian else INPUT_FILL
                if index in (1, 4):
                    cell.font = Font(name=FONT, size=9, color="808080")

        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = f"A1:D{len(rows) + 1}"
        for column, width in zip("ABCD", (38, 60, 60, 46)):
            sheet.column_dimensions[column].width = width

        counts[namespace] = len(rows)

    # --- Overview -----------------------------------------------------------
    overview.column_dimensions["A"].width = 22
    for column in "BCD":
        overview.column_dimensions[column].width = 14

    title = write_cell(overview, 1, 1, "Georgian translation")
    title.font = Font(name=FONT, bold=True, size=16)

    legend = [
        "",
        "Fill in the GEORGIAN column (shaded) on each tab. Leave key, english and source alone.",
        "Rows still to do are sorted to the top of every tab. Green means already translated.",
        "",
        "Example of a finished row, from the wallet tab:",
        "    key       savedCardsPanel.savedCards",
        "    english   Saved cards",
        "    georgian  შენახული ბარათები",
        "",
        "Keep {placeholders} in braces exactly as they are — word order around them can change.",
        "Do not translate: the brand name, IDs, licence plates, ₾, API, SMS.",
        "Formal address (თქვენ). Buttons are verbal nouns — შენახვა, not შეინახე.",
        "Read src/messages/GLOSSARY.md before starting: it settles Order / Load / Job / Cargo,",
        "which are four different words in Georgian and must stay different.",
        "",
        "When done, save this file and run:  python3 scripts/translations_xlsx.py --merge",
    ]
    for offset, line in enumerate(legend, start=2):
        cell = write_cell(overview, offset, 1, line)
        cell.font = Font(name=FONT, bold=line.startswith("Example"))

    header_row = len(legend) + 3
    for index, name in enumerate(["Tab", "Strings", "Done", "Remaining"], start=1):
        cell = write_cell(overview, header_row, index, name)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL

    for offset, namespace in enumerate(namespaces(), start=header_row + 1):
        last = counts[namespace] + 1
        write_cell(overview, offset, 1, namespace).font = Font(name=FONT)
        write_cell(overview, offset, 2, counts[namespace]).font = Font(name=FONT)
        # Live counts, so the workbook reports its own progress as it is filled.
        overview.cell(
            row=offset,
            column=3,
            value=f"=COUNTA('{namespace}'!C2:C{last})",
        ).font = Font(name=FONT)
        overview.cell(
            row=offset, column=4, value=f"=B{offset}-C{offset}"
        ).font = Font(name=FONT)

    total_row = header_row + len(counts) + 1
    for index, formula in enumerate(
        [
            "TOTAL",
            f"=SUM(B{header_row + 1}:B{total_row - 1})",
            f"=SUM(C{header_row + 1}:C{total_row - 1})",
            f"=SUM(D{header_row + 1}:D{total_row - 1})",
        ],
        start=1,
    ):
        cell = overview.cell(row=total_row, column=index, value=formula)
        cell.font = Font(name=FONT, bold=True)

    overview.freeze_panes = f"A{header_row + 1}"
    book.save(WORKBOOK)

    print(f"{sum(counts.values())} strings across {len(counts)} tabs")
    print(f"-> {WORKBOOK.relative_to(ROOT)}")


def merge():
    if not WORKBOOK.exists():
        sys.exit(f"{WORKBOOK.relative_to(ROOT)} not found — build it first.")

    book = load_workbook(WORKBOOK, data_only=True)
    merged = 0
    unknown = 0

    for namespace in namespaces():
        if namespace not in book.sheetnames:
            continue

        entries = load(namespace)
        sheet = book[namespace]
        header = [c.value for c in sheet[1]]

        try:
            key_at = header.index("key")
            georgian_at = header.index("georgian")
        except ValueError:
            print(f"{namespace}: header row was renamed — skipped.")
            continue

        count = 0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            key = (row[key_at] or "").strip() if row[key_at] else ""
            georgian = str(row[georgian_at]).strip() if row[georgian_at] else ""
            if not key or not georgian:
                continue
            if key not in entries:
                unknown += 1
                continue
            if entries[key].get("ka") != georgian:
                count += 1
            entries[key]["ka"] = georgian

        (JSON_DIR / f"{namespace}.json").write_text(
            json.dumps(entries, indent=2, ensure_ascii=False) + "\n",
            encoding="utf-8",
        )
        if count:
            print(f"{namespace:<11} {count} updated")
        merged += count

    print(f"\n{merged} translations merged into translation/*.json")
    if unknown:
        print(f"{unknown} row(s) had keys no longer in the source — re-export those.")
    print("Next: pnpm i18n:import")


if __name__ == "__main__":
    merge() if "--merge" in sys.argv else build()
